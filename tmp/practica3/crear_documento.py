from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'practica3'
doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.space_before = Pt(0)
    if name != 'Normal':
        style.font.bold = True
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + key), None)
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + key), 'Arial')

def p(text):
    return doc.add_paragraph(text)

def heading(text):
    doc.add_paragraph(text, 'Heading 1')

def picture(name, width=7.1):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(7)
    shape = paragraph.add_run().add_picture(str(OUT / 'capturas' / name), width=Inches(width))
    shape._inline.docPr.set('descr', name.replace('_', ' ').replace('.png', ''))
    if name.startswith(('01_', '02_')):
        fill = shape._inline.graphic.graphicData.pic.blipFill
        crop = OxmlElement('a:srcRect')
        crop.set('r', '18000')
        crop.set('b', '16000')
        fill.insert(1, crop)

doc.add_paragraph('Práctica 3 Categorías y relaciones', 'Title')
p('En esta práctica se relacionaron los productos con sus categorías en el mismo proyecto Django. Se comprobó que la API muestra el nombre de cada categoría y permite consultar solo los productos de una categoría elegida.')
heading('Categorías en el administrador')
p('Se registraron Electrónica con el ID 1 y Ropa con el ID 2. El administrador muestra las dos categorías guardadas.')
picture('01_admin_categorias.png')
heading('Productos con su categoría asignada')
p('Mouse y Teclado pertenecen a Electrónica, mientras que Playera pertenece a Ropa. La columna Categoria permite revisar esta relación en el administrador.')
picture('02_admin_productos.png')

doc.add_page_break()
heading('Consulta de categorías en la API')
p('La dirección /api/categories/ devuelve las dos categorías registradas. Cada una aparece con su identificador y su nombre, y la respuesta HTTP 200 OK indica que la consulta funcionó.')
picture('03_api_categorias.png')
heading('Relación entre categorías y productos')
p('El modelo Category guarda el nombre de la categoría. El modelo Product utiliza una ForeignKey para relacionar cada producto con una categoría, de modo que varios productos pueden pertenecer a la misma.')
p('Se dejó la categoría como un dato obligatorio y se aplicó la migración correspondiente. También se registraron ambos modelos en el administrador y se agregó la dirección /api/ para realizar las consultas solicitadas.')

doc.add_page_break()
heading('Productos con el nombre de su categoría')
p('La dirección /api/products/ muestra los tres productos con su nombre, descripción, precio y categoría. En el campo category aparece Electrónica o Ropa, en lugar de mostrar solamente un número.')
picture('04_api_productos.png')
heading('Cambio realizado para mostrar el nombre')
p('En ProductSerializer se agregó StringRelatedField al campo category para mostrar el nombre de la categoría. Para guardar o cambiar un producto desde la API, el campo category_id recibe el identificador de una categoría existente.')

doc.add_page_break()
heading('Filtro de productos por categoría')
p('La consulta /api/products/?category=1 muestra únicamente Mouse y Teclado, que pertenecen a Electrónica. Playera ya no aparece porque está relacionada con Ropa, cuyo identificador es 2.')
picture('05_api_filtro.png')
heading('Configuración del filtro')
p('Se agregó django_filters a la configuración del proyecto y DjangoFilterBackend a la vista de productos. Con filterset_fields se indicó que el campo category puede usarse para filtrar la lista.')
heading('Resultado de la práctica')
p('Las cinco capturas muestran las categorías, la relación con los productos y el filtro funcionando. Además, se realizaron siete pruebas automáticas que terminaron sin errores, incluyendo la consulta de categorías, el nombre visible y la separación de productos por categoría.')

doc.core_properties.title = 'Práctica 3 Categorías y relaciones'
doc.core_properties.subject = 'Categorías, productos y filtrado en Django'
doc.core_properties.author = ''
doc.core_properties.keywords = ''
doc.save(OUT / 'Practica_3_Categorias_y_Relaciones.docx')
print(OUT / 'Practica_3_Categorias_y_Relaciones.docx')
