from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parents[1] / 'output' / 'unidad1'
OUT.mkdir(parents=True, exist_ok=True)

for name, box in {
    'api': (0, 0, 850, 490),
    'models': (0, 0, 720, 305),
    'tests': (0, 0, 720, 465),
    'terminal': (0, 0, 720, 330),
    'services': (0, 0, 850, 910),
}.items():
    im = Image.open(ROOT / f'{name}.png')
    im.crop(box).save(ROOT / f'{name}_recorte.png')

doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.25)
section.right_margin = Cm(2.25)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(12)
normal.font.color.rgb = None
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for sname, size in [('Title', 12), ('Heading 1', 12), ('Heading 2', 12)]:
    style = styles[sname]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = None
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(7)

def add_p(text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def heading(text):
    doc.add_heading(text, level=1)

def caption(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Arial'
    r.font.size = Pt(12)
    return p

def picture(name, width_cm=16.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(ROOT / f'{name}_recorte.png'), width=Cm(width_cm))

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '6')
        edge.set(qn('w:color'), '000000')
        borders.append(edge)
    tcPr.append(borders)

def set_cell_text(cell, value, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(value)
    run.bold = bold
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run('Resultado de aprendizaje — Unidad 1')
title_run.bold = True
title_run.font.name = 'Arial'
title_run.font.size = Pt(12)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run('Proyecto Django REST: precios finales por categoría').bold = True

heading('Arquitectura del proyecto')
add_p('El proyecto usa la estructura MVT de Django y está organizado en capas. La ruta recibe la petición, la vista revisa el acceso con JWT, el servicio calcula el precio y el modelo obtiene los datos del producto. Así, el descuento queda en un solo lugar y la vista es más fácil de entender.')

flow = doc.add_table(rows=1, cols=7)
flow.alignment = WD_TABLE_ALIGNMENT.CENTER
flow.autofit = False
for cell, label in zip(flow.rows[0].cells,
                       ['urls.py', '→', 'views.py', '→', 'services.py', '→', 'models.py']):
    set_cell_text(cell, label, bold=(label != '→'))
for i, width in enumerate([2.7, .5, 2.7, .5, 3.6, .5, 3.0]):
    flow.columns[i].width = Cm(width)
    flow.cell(0, i).width = Cm(width)
doc.add_paragraph()
add_p('El recorrido empieza en /api/products/<id>/precio-final/. La vista pide el resultado al servicio y el servicio consulta el producto y su categoría; después regresa el precio final en formato JSON.')

doc.add_page_break()
heading('Strategy: descuento según la categoría')
add_p('En products/services.py se definieron tres formas de calcular el precio: 10 % de descuento para Electrónica, 20 % para Ropa y sin descuento para las demás categorías. El servicio elige la opción que corresponde al nombre de la categoría y calcula con dos decimales.')
caption('Captura 1. Estrategias de descuento en services.py')
picture('services', 16.4)

doc.add_page_break()
heading('Factory Method: categoría predeterminada')
add_p('El método crear_con_categoria_default crea un producto cuando no se indica una categoría. Primero busca o crea “Sin categoría” y luego guarda el producto con esa relación. Esto evita repetir esos pasos cada vez que se necesite un producto así.')
caption('Captura 2. Método en models.py')
picture('models')

heading('Facade: consulta del precio final')
add_p('El endpoint /api/products/<id>/precio-final/ está protegido con JWT. La vista hace una sola llamada al servicio y entrega el resultado; por eso funciona como una entrada sencilla al cálculo completo. La ruta también está disponible bajo /api/v1/.')

doc.add_page_break()
heading('Prueba con dos categorías')
add_p('Se consultó el endpoint con un token de acceso para dos productos reales del proyecto. Mouse cuesta 250.00 y, al ser de Electrónica, queda en 225.00; Playera cuesta 180.00 y, al ser de Ropa, queda en 144.00. Las dos peticiones respondieron con código 200.')
caption('Captura 3. Respuestas del endpoint con JWT')
picture('api')

doc.add_page_break()
heading('Pruebas del endpoint')
add_p('Se agregaron dos pruebas nuevas. La primera consulta dos productos con acceso autorizado y comprueba el precio de cada uno; la segunda consulta sin token y espera el código 401. Las pruebas anteriores también siguen funcionando.')
caption('Captura 4. Pruebas nuevas en tests.py')
picture('tests')
caption('Captura 5. Salida de python manage.py test')
picture('terminal')

doc.add_page_break()
heading('Tabla de patrones de diseño')
add_p('Los tres patrones se aplicaron en partes distintas del proyecto. Cada uno ayuda a ordenar una tarea concreta: crear productos, elegir descuentos o responder desde la API.')

rows = [
    ('Factory Method', 'Creacional', 'Método de Product en models.py',
     'Crea el producto en “Sin categoría”.'),
    ('Strategy', 'Comportamiento', 'Clases de descuento en services.py',
     'Elige 10 % o 20 % según la categoría.'),
    ('Facade', 'Estructural', 'precio_final() en views.py',
     'Da una sola entrada al cálculo del servicio.'),
]
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
headers = ['Nombre', 'Categoría', 'Evidencia', 'Explicación']
for cell, value in zip(table.rows[0].cells, headers):
    set_cell_text(cell, value, bold=True)
for row in rows:
    cells = table.add_row().cells
    for cell, value in zip(cells, row):
        cell.text = value
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(12)
        border(cell)
for i, width in enumerate([3.0, 2.8, 5.1, 5.4]):
    table.columns[i].width = Cm(width)

output = OUT / 'Resultado_aprendizaje_Unidad1.docx'
doc.save(output)
print(output)
