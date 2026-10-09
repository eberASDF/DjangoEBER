from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'practica4'
CAPTURAS = OUT / 'capturas'
DOCX = OUT / 'Practica_4_Autenticacion_JWT.docx'

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
    style = doc.styles[style_name]
    style.font.name = 'Arial'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.space_before = Pt(0)
    if style_name != 'Normal':
        style.font.bold = True
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + key), None)
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + key), 'Arial')


def paragraph(text):
    item = doc.add_paragraph(text)
    item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return item


def heading(text):
    item = doc.add_paragraph(text, 'Heading 1')
    item.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return item


def picture(filename, width=7.1, crop_right=0, crop_bottom=0):
    item = doc.add_paragraph()
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER
    item.paragraph_format.space_after = Pt(6)
    shape = item.add_run().add_picture(str(CAPTURAS / filename), width=Inches(width))
    shape._inline.docPr.set('descr', filename.replace('_', ' ').replace('.png', ''))
    if crop_right or crop_bottom:
        fill = shape._inline.graphic.graphicData.pic.blipFill
        crop = OxmlElement('a:srcRect')
        if crop_right:
            crop.set('r', str(crop_right))
        if crop_bottom:
            crop.set('b', str(crop_bottom))
        fill.insert(1, crop)


title = doc.add_paragraph('Práctica 4 Autenticación con JWT', 'Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
paragraph(
    'En esta práctica se agregó autenticación JWT a la API del proyecto Django. '
    'Ahora las categorías y los productos están protegidos, por lo que se necesita '
    'un token válido para consultar o modificar la información.'
)

heading('Paso 1 Instalar Simple JWT')
paragraph(
    'Se instaló djangorestframework-simplejwt dentro del entorno virtual del proyecto. '
    'La instalación agregó la versión 5.5.1 de Simple JWT y la biblioteca PyJWT 2.15.0.'
)
picture('01_instalacion.png')

doc.add_page_break()
heading('Paso 2 Configurar settings.py')
paragraph(
    'Se agregó rest_framework_simplejwt a las aplicaciones instaladas. También se '
    'configuró JWTAuthentication como el método que Django REST Framework usa para '
    'revisar la identidad del usuario.'
)
picture('02_settings.png', width=6.9)

doc.add_page_break()
heading('Paso 3 Proteger los endpoints')
paragraph(
    'En products/views.py se importó IsAuthenticated y se agregó permission_classes '
    'a ProductViewSet y CategoryViewSet. Con este cambio, los productos y las '
    'categorías solo responden cuando la solicitud incluye una autenticación válida.'
)
picture('03_views.png', width=6.9)

doc.add_page_break()
heading('Paso 4 Agregar las rutas de token')
paragraph(
    'En tienda_online/urls.py se agregaron las rutas /api/token/ y '
    '/api/token/refresh/. La primera entrega los tokens access y refresh, mientras '
    'que la segunda permite obtener un nuevo token de acceso.'
)
picture('04_urls.png', width=6.9)

doc.add_page_break()
heading('Comprobación sin token')
paragraph(
    'Antes de obtener un token se consultó /api/products/. La respuesta fue HTTP 401 '
    'Unauthorized, lo que demuestra que el endpoint ya no permite el acceso anónimo.'
)
picture('05_sin_token_401.png', width=6.9, crop_bottom=38000)

doc.add_page_break()
heading('Paso 5 Obtener el token')
paragraph(
    'Se enviaron el usuario y la contraseña de una cuenta temporal a /api/token/. '
    'La respuesta HTTP 200 OK incluyó un token refresh y un token access; en la '
    'captura solo se muestran partes de cada token para no exponerlos completos.'
)
picture('06_token_obtenido.png', width=6.8, crop_right=19000)

heading('Uso del token')
paragraph(
    'El token access debe enviarse en el encabezado Authorization con la palabra '
    'Bearer. Este encabezado permite que Django reconozca al usuario durante cada '
    'consulta protegida.'
)

doc.add_page_break()
heading('Acceso a productos con el token')
paragraph(
    'Se consultó /api/products/ enviando el token en Authorization. La API respondió '
    'HTTP 200 OK y mostró los productos Mouse, Playera y Teclado con sus categorías, '
    'por lo que la autenticación funcionó correctamente.'
)
picture('07_productos_con_token.png', width=6.8)

heading('Resultado de la práctica')
paragraph(
    'La API quedó protegida con JWT y las rutas para obtener y renovar tokens están '
    'disponibles. También se ejecutaron diez pruebas automáticas sin errores, '
    'incluyendo el rechazo sin token, la obtención del token, su renovación y el '
    'acceso autorizado a los productos.'
)

doc.core_properties.title = 'Práctica 4 Autenticación con JWT'
doc.core_properties.subject = 'Autenticación JWT en una API de Django'
doc.core_properties.author = ''
doc.core_properties.keywords = ''
doc.save(DOCX)
print(DOCX)
